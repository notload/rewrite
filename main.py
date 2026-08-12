import openpyxl
import random
import os
from pathlib import Path
from datetime import datetime
import requests
import tempfile
import time
import subprocess
from kivy.app import App
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.textinput import TextInput
from kivy.uix.button import Button
from kivy.uix.label import Label
from kivy.uix.scrollview import ScrollView
from kivy.clock import Clock
from kivy.utils import platform
from kivy.core.audio import SoundLoader

class DictationApp(BoxLayout):
    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.orientation = 'vertical'
        
        self.output = Label(text="欢迎使用听写程序\n请先加载 Excel 文件", 
                            size_hint_y=None, height=400, 
                            text_size=(400, None), halign='left', valign='top')
        scroll = ScrollView(size_hint=(1, 0.8))
        scroll.add_widget(self.output)
        self.add_widget(scroll)
        
        input_box = BoxLayout(size_hint=(1, 0.1))
        self.input = TextInput(hint_text="输入答案或指令", multiline=False)
        send_btn = Button(text="发送", size_hint_x=0.2)
        send_btn.bind(on_press=self.on_send)
        input_box.add_widget(self.input)
        input_box.add_widget(send_btn)
        self.add_widget(input_box)
        
        self.status = Label(text="状态：等待加载", size_hint=(1, 0.1))
        self.add_widget(self.status)
        
        self.rows = []           # 所有行数据，每行为 [值1, 值2, 错误次数]
        self.mode = None
        self.col_pick = None
        self.col_show = None
        self.col_update = None
        self.extract_counts = {}
        self.answer_counts = {}
        self.total_questions = 0
        self.correct_answers = 0
        self.idx = None
        self.row = None
        self.waiting_for_answer = False
        self.user_answer = None
        self.file_path = None
        
        self.append_output("请先输入 'load 文件路径' 来加载 Excel")
        self.append_output("示例：load /sdcard/听写文件.xlsx")
    
    def append_output(self, text):
        self.output.text = self.output.text + "\n" + text
    
    def on_send(self, instance):
        cmd = self.input.text.strip()
        self.input.text = ""
        if not cmd:
            return
        self.process_command(cmd)
    
    def process_command(self, cmd):
        if cmd.lower().startswith("load "):
            path = cmd[5:].strip()
            self.load_file(path)
        elif cmd.lower() == "stop":
            self.save_and_exit()
        elif cmd.lower().startswith("mode "):
            try:
                mode = int(cmd[5:])
                self.set_mode(mode)
            except:
                self.append_output("格式错误，输入 'mode 1' 或 'mode 2'")
        elif cmd.lower() == "play":
            if self.row is not None:
                self.play_audio(str(self.row[self.col_show]))
            else:
                self.append_output("无内容可播放")
            if hasattr(self, 'audio_pending') and self.audio_pending:
                self.audio_pending = False
                self.total_questions += 1
                self.status.text = f"已完成 {self.total_questions} 题"
                Clock.schedule_once(lambda dt: self.next_question(), 0.1)
        elif cmd.lower() in ('y', 'n'):
            if hasattr(self, 'answer_pending') and self.answer_pending:
                self.answer_pending = False
                self.handle_judgement(cmd)
            else:
                self.append_output("当前没有等待确认")
        elif self.waiting_for_answer:
            self.user_answer = cmd
            self.waiting_for_answer = False
            self.continue_question()
        else:
            self.append_output("未知指令")
    
    def load_file(self, path):
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            sheet = wb.active
            self.rows = []
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    # 确保有3列，缺的补0
                    while len(row) < 3:
                        row = list(row) + [0]
                    self.rows.append(list(row))
            self.append_output(f"成功加载文件：{path}，共 {len(self.rows)} 行")
            self.file_path = path
            self.append_output("请选择模式：输入 'mode 1' 或 'mode 2'")
            self.status.text = "请选择模式"
        except Exception as e:
            self.append_output(f"加载失败：{e}")
    
    def set_mode(self, mode):
        if mode == 1:
            self.col_pick, self.col_show, self.col_update = 0, 1, 2
        elif mode == 2:
            self.col_pick, self.col_show, self.col_update = 1, 0, 3
            # 确保每行至少有4列
            for row in self.rows:
                while len(row) < 4:
                    row.append(0)
        else:
            self.append_output("无效模式，请输入 1 或 2")
            return
        self.mode = mode
        self.append_output(f"模式已设为 {mode}")
        self.status.text = "开始听写..."
        self.next_question()
    
    def next_question(self):
        if not self.rows:
            self.append_output("请先加载文件")
            return
        
        idx = self.weighted_choice()
        self.idx = idx
        self.row = self.rows[idx]
        
        self.extract_counts[idx] = self.extract_counts.get(idx, 0) + 1
        extract_count = self.extract_counts[idx]
        error_count = int(self.row[self.col_update])
        answer_count = self.answer_counts.get(idx, 0)
        error_rate = (error_count / answer_count * 100) if answer_count > 0 else 0.0
        
        self.append_output(f"\n--- 第 {self.total_questions+1} 题 ---")
        self.append_output(f"词条：{self.row[self.col_pick]}")
        self.append_output(f"总抽取：{extract_count}，已作答：{answer_count}，错误：{error_count}，错误率：{error_rate:.1f}%")
        if error_count >= 5:
            self.append_output("⚠️ 错误较多，重点复习！")
        elif error_count >= 3:
            self.append_output("⚡ 错误较多，认真作答！")
        
        self.append_output("请输入你的答案（或输入 'skip' 查看答案）")
        self.waiting_for_answer = True
        self.status.text = "等待输入答案"
    
    def weighted_choice(self):
        weights = [(row[self.col_update] + 1) ** 2 for row in self.rows]
        if all(w == weights[0] for w in weights):
            return random.randint(0, len(self.rows) - 1)
        return random.choices(range(len(self.rows)), weights=weights, k=1)[0]
    
    def continue_question(self):
        user_answer = self.user_answer.strip()
        self.user_answer = None
        
        if user_answer.lower() == "skip":
            self.append_output(f"正确内容：{self.row[self.col_show]}")
            self.append_output("如需读音，输入 'play'")
            self.waiting_for_answer = True
            return
        
        correct = str(self.row[self.col_show]).strip().lower()
        if user_answer and user_answer.lower() == correct:
            self.append_output("✅ 自动匹配正确！")
            answer = 'y'
        else:
            self.append_output(f"你的答案：{user_answer}，正确内容：{self.row[self.col_show]}")
            self.append_output("输入 'y' 表示正确，'n' 表示错误")
            self.waiting_for_answer = True
            self.answer_pending = True
            return
        
        self.handle_judgement(answer)
    
    def handle_judgement(self, answer):
        idx = self.idx
        self.answer_counts[idx] = self.answer_counts.get(idx, 0) + 1
        
        if answer.lower() == 'y':
            self.correct_answers += 1
            self.append_output("✅ 回答正确！")
        else:
            self.rows[idx][self.col_update] += 1
            new_error = self.rows[idx][self.col_update]
            self.append_output(f"❌ 错误已记录，当前错误次数：{new_error}")
            self.append_output("播放读音？输入 'play' 或 'no'")
            self.waiting_for_answer = True
            self.audio_pending = True
            return
        
        self.total_questions += 1
        self.status.text = f"已完成 {self.total_questions} 题"
        Clock.schedule_once(lambda dt: self.next_question(), 0.1)
    
    def play_audio(self, text):
        lang = "zh" if any('\u4e00' <= ch <= '\u9fff' for ch in text) else "en"
        url = "http://fanyi.baidu.com/gettts"
        params = {"lan": lang, "text": text, "spd": 3, "source": "web"}
        try:
            resp = requests.get(url, params=params, timeout=5)
            if resp.status_code == 200:
                with tempfile.NamedTemporaryFile(delete=False, suffix=".mp3") as f:
                    f.write(resp.content)
                    tmp_path = f.name
                if platform == 'android':
                    sound = SoundLoader.load(tmp_path)
                    if sound:
                        sound.play()
                        time.sleep(sound.length + 0.5)
                    else:
                        self.append_output("播放失败")
                else:
                    subprocess.run(["ffplay", "-nodisp", "-autoexit", tmp_path], check=False)
                os.unlink(tmp_path)
            else:
                self.append_output("获取语音失败")
        except Exception as e:
            self.append_output(f"播放出错：{e}")
    
    def save_and_exit(self):
        if not self.rows:
            self.append_output("没有数据可保存")
            return
        now = datetime.now()
        month = f"{now.month:02d}"
        day = f"{now.day:02d}"
        prefix = f"{month}_{day}"
        save_dir = Path("/sdcard/听写记录")
        save_dir.mkdir(parents=True, exist_ok=True)
        max_num = 0
        for f in save_dir.glob(f"{prefix}_*.xlsx"):
            stem = f.stem
            parts = stem.split('_')
            if len(parts) == 3 and parts[2].isdigit():
                num = int(parts[2])
                if num > max_num:
                    max_num = num
        next_num = max_num + 1
        filename = f"{prefix}_{next_num}.xlsx"
        save_path = save_dir / filename
        
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            for row in self.rows:
                ws.append(row)
            wb.save(save_path)
            self.append_output(f"✅ 已保存至：{save_path}")
        except Exception as e:
            self.append_output(f"保存失败：{e}")
        App.get_running_app().stop()

class DictationAppMain(App):
    def build(self):
        return DictationApp()

if __name__ == '__main__':
    DictationAppMain().run()
